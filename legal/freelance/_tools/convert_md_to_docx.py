"""Convertit les .md du pack juridique freelance en .docx professionnels.

Usage : python .convert_md_to_docx.py

Sources : *.md du dossier courant
Sortie : word/*.docx
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).parent
OUT = HERE / 'word'
OUT.mkdir(exist_ok=True)


def set_cell_shading(cell, fill_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill_color)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def add_inline_runs(paragraph, text: str):
    """Parse **bold**, *italic*, [link](url) dans le texte et ajoute les runs."""
    pattern = re.compile(
        r'(\*\*([^*]+)\*\*)'                        # bold
        r'|(\*([^*]+)\*)'                           # italic
        r'|(\[([^\]]+)\]\(([^)]+)\))'               # link
        r'|(`([^`]+)`)',                            # code
    )
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group(1):  # bold
            r = paragraph.add_run(m.group(2))
            r.bold = True
        elif m.group(3):  # italic
            r = paragraph.add_run(m.group(4))
            r.italic = True
        elif m.group(5):  # link
            r = paragraph.add_run(m.group(6))
            r.italic = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif m.group(8):  # code
            r = paragraph.add_run(m.group(9))
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_paragraph_with_inline(doc, text: str, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    add_inline_runs(p, text.strip())
    return p


def parse_table(lines, start_idx):
    """Parse un tableau markdown commencant a lines[start_idx]. Retourne (rows, end_idx)."""
    rows = []
    i = start_idx
    while i < len(lines) and lines[i].strip().startswith('|'):
        line = lines[i].strip()
        if line.startswith('|') and re.match(r'^\|[\s:|-]+\|?$', line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows.append(cells)
        i += 1
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_cells in enumerate(rows):
        for c_idx in range(cols):
            cell = table.rows[r_idx].cells[c_idx]
            text = row_cells[c_idx] if c_idx < len(row_cells) else ''
            cell.text = ''
            p = cell.paragraphs[0]
            add_inline_runs(p, text)
            if r_idx == 0:
                set_cell_shading(cell, 'D5E8F0')
                for run in p.runs:
                    run.bold = True


def convert_md_to_docx(md_path: Path, docx_path: Path, title: str):
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    raw = md_path.read_text(encoding='utf-8')
    lines = raw.split('\n')
    i = 0
    in_blockquote = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            doc.add_paragraph()
            i += 1
            continue

        if stripped == '---':
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '999999')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        if stripped.startswith('# '):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith('### '):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:].strip(), level=3)
            i += 1
            continue

        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|?$', lines[i+1].strip()):
            rows, new_i = parse_table(lines, i)
            add_table(doc, rows)
            i = new_i
            continue

        if stripped.startswith('> '):
            content = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run('« ')
            r.italic = True
            r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            add_inline_runs(p, content)
            for run in p.runs[1:]:
                run.italic = True
                if run.font.color.rgb is None:
                    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            p.add_run(' »').italic = True
            i += 1
            continue

        if stripped.startswith('- ') or stripped.startswith('* '):
            content = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            add_inline_runs(p, content)
            i += 1
            continue

        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            content = m.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            add_inline_runs(p, content)
            i += 1
            continue

        if stripped.startswith('- [ ]') or stripped.startswith('- [x]'):
            done = stripped.startswith('- [x]')
            content = stripped[5:].strip()
            p = doc.add_paragraph()
            r = p.add_run('☑ ' if done else '☐ ')
            add_inline_runs(p, content)
            i += 1
            continue

        m_alpha = re.match(r'^([a-z])\)\s+(.*)$', stripped)
        if m_alpha:
            content = m_alpha.group(2).strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            r = p.add_run(f"{m_alpha.group(1)}) ")
            r.bold = True
            add_inline_runs(p, content)
            i += 1
            continue

        add_paragraph_with_inline(doc, stripped)
        i += 1

    doc.save(docx_path)


def main():
    files = sorted([p for p in HERE.glob('*.md') if not p.name.startswith('.')])
    print(f"Conversion de {len(files)} fichiers...")
    for md in files:
        out = OUT / (md.stem + '.docx')
        title = md.stem.replace('_', ' ')
        try:
            convert_md_to_docx(md, out, title)
            print(f"  OK : {md.name} -> word/{out.name}")
        except Exception as e:
            print(f"  KO : {md.name} -> {e}")
            raise


if __name__ == '__main__':
    main()
